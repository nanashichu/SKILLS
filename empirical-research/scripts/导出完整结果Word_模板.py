"""从指定Word继承版式，填入已核实CSV与图片，不估计或改动统计结果。

python 导出完整结果Word_模板.py manifest.json --output 测试/初稿整理/20261010_主题/20261010_完整结果.docx --full-results --project-root 项目根
清单必须提供reference_docx；每个表提供reference_table（从1计数）。
方法来自知识产权法庭项目定稿/code/04_核验并导出剔除京沪及广东非深圳_20261006.py。
清单格式见 references/report-generation.md。依赖python-docx。
"""
import argparse
import csv
import hashlib
import json
import re
import shutil
from datetime import datetime
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table, _Cell
from lxml import etree


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def require(condition, message):
    if not condition:
        raise ValueError(message)


def validate_output_name(output):
    match = re.fullmatch(r'(\d{8})_(.+)\.docx', output.name, re.I)
    require(match is not None, 'Word文件名必须为YYYYMMDD_主题.docx，日期在前')
    try:
        datetime.strptime(match[1], '%Y%m%d')
    except ValueError:
        raise ValueError('Word文件名的前8位必须是有效日期') from None


def copy_complete_word(output, project_root):
    """完整汇编额外保留阅读副本；不覆盖不同内容的同名文件。"""
    validate_output_name(output)
    require(project_root.is_dir(), '完整汇编必须指定已存在的真实项目根目录')
    target = project_root / '完整实证' / output.name
    source_hash = sha(output)
    target.parent.mkdir(parents=True, exist_ok=True)
    if target.exists():
        require(sha(target) == source_hash, '完整实证已有不同内容的同名Word，请先保留可恢复归档')
    else:
        with output.open('rb') as src, target.open('xb') as dst:
            shutil.copyfileobj(src, dst)
    require(sha(output) == sha(target) == source_hash, '完整实证副本与来源Word的SHA256不一致')
    return target


def xml(node):
    return None if node is None else etree.tostring(node, method='c14n', exclusive=True)


def read_rows(path):
    with path.open(encoding='utf-8-sig', newline='') as f:
        rows = list(csv.reader(f))
    require(rows and len(rows[0]) >= 2, f'CSV表格为空或不足两列: {path}')
    require(all(len(row) == len(rows[0]) for row in rows), f'CSV列数不一致: {path}')
    return rows


def table_skeleton(source, rows):
    """保留原三线边框和列宽；允许像项目补R²一样增加最后一行。"""
    element = deepcopy(source._tbl)
    original = element.findall(qn('w:tr'))
    require(len(rows) in (len(original), len(original) + 1),
            'CSV与参考表行数不匹配；仅支持同样行数或增加一行统计量')
    if len(rows) == len(original) + 1:
        last = original[-1]
        added = deepcopy(last)
        for border in last.xpath('./w:tc/w:tcPr/w:tcBorders/w:bottom'):
            border.set(qn('w:val'), 'nil')
        element.append(added)
    for tr, values in zip(element.findall(qn('w:tr')), rows):
        require(len(tr.findall(qn('w:tc'))) == len(values),
                'CSV与参考表单元格数不匹配；先明确合并单元格映射')
    return element


def set_cell_text(cell, value):
    """沿用参考单元格的首段/首文字格式，不重新规定字号和字体。"""
    p = cell.paragraphs[0]
    ppr = deepcopy(p._p.pPr)
    rpr = deepcopy(p.runs[0]._r.rPr) if p.runs else None
    cell.text = value
    for paragraph in cell.paragraphs:
        if paragraph._p.pPr is not None:
            paragraph._p.remove(paragraph._p.pPr)
        if ppr is not None:
            paragraph._p.insert(0, deepcopy(ppr))
        for run in paragraph.runs:
            if run._r.rPr is not None:
                run._r.remove(run._r.rPr)
            if rpr is not None:
                run._r.insert(0, deepcopy(rpr))


def verify(docx, tables, source, figures):
    doc = Document(docx)
    require(len(doc.sections) == 1, '输出分节数改变')
    require(xml(doc.sections[0]._sectPr) == xml(source.sections[0]._sectPr),
            '页面设置与参考Word不一致')
    require(xml(doc.styles.element) == xml(source.styles.element), '样式与参考Word不一致')
    require(len(doc.tables) == len(tables), 'Word表数不一致')
    for table, item in zip(doc.tables, tables):
        expected = item['filled_element']
        require(xml(table._tbl) == xml(expected), '表格格式或内容改变: ' + item['title'])
        values = [[_Cell(tc, table).text for tc in tr.findall(qn('w:tc'))]
                  for tr in table._tbl.findall(qn('w:tr'))]
        require(values == item['rows'], 'Word与CSV不同: ' + item['title'])
    require(len(doc.inline_shapes) == figures, '图片数量错误')
    return doc


def build(manifest, output, *, full_results=False, project_root=None):
    require(output.suffix.lower() == '.docx', 'Word输出文件必须使用.docx后缀')
    validate_output_name(output)
    if full_results:
        require(project_root is not None and project_root.is_dir(),
                '全部实证结果Word必须传入--project-root，额外副本保存到完整实证')
        require(not output.resolve().is_relative_to((project_root / '完整实证').resolve()),
                '完整实证只放阅读副本；原Word及回执请输出到测试任务目录')
    require(all('初稿' not in p.parts for p in (output.absolute(), output.resolve())),
            '初稿由用户放置论文，结果文档请输出到测试任务目录')
    require(not output.exists() and not output.with_suffix('.核验.json').exists(),
            '输出Word或核验回执已存在，请先保留可恢复归档，再导出到同一任务路径')
    manifest_hash = sha(manifest)
    config = json.loads(manifest.read_text(encoding='utf-8-sig'))
    require(isinstance(config.get('reference_docx'), str) and config['reference_docx'],
            '必须指定reference_docx，以现用稿件为版式来源')
    reference = (manifest.parent / config['reference_docx']).resolve()
    reference_hash = sha(reference)
    source = Document(reference)
    require(len(source.sections) == 1, '本模板适用于项目的一节文档；多节稿件需按实际节结构适配')
    paragraphs = config.get('paragraphs', [])
    require(isinstance(paragraphs, list) and all(isinstance(p, str) for p in paragraphs),
            'paragraphs须为字符串列表')
    prepared = []
    for block in config['blocks']:
        item = dict(block)
        require(isinstance(item.get('notes', []), list)
                and all(isinstance(n, str) for n in item.get('notes', [])), 'notes须为字符串列表')
        if item['type'] in ('table', 'figure'):
            path = (manifest.parent / item['path']).resolve()
            if not path.is_file():
                raise FileNotFoundError(path)
            item['source_path'], item['sha256'] = path, sha(path)
            if item['type'] == 'table':
                number = item.get('reference_table')
                require(type(number) is int and 1 <= number <= len(source.tables),
                        'reference_table必须是参考Word中从1计数的表序号')
                item['rows'] = read_rows(path)
                item['element'] = table_skeleton(source.tables[number - 1], item['rows'])
        elif item['type'] != 'paragraph':
            raise ValueError(f"不支持的块类型: {item['type']}")
        prepared.append(item)
    inputs = {manifest.resolve(), reference} | {x['source_path'] for x in prepared if 'source_path' in x}
    outputs = {output.resolve(), output.with_suffix('.核验.json').resolve()}
    require(not inputs & outputs, 'Word或回执输出与来源路径冲突，禁止覆盖输入')
    # 从指定稿件保留页面、样式和设置，清除正文后按清单重建。
    doc = Document(reference)
    # OOXML要求zoom有percent；修复来源中缺失的显示属性，不改变页面/表格版式。
    compatibility_fixes = []
    zoom = doc.settings.element.find(qn('w:zoom'))
    if zoom is not None and zoom.get(qn('w:percent')) is None:
        zoom.set(qn('w:percent'), '100')
        compatibility_fixes.append('补齐zoom.percent=100以通过OOXML校验')
    for child in list(doc._element.body):
        if child.tag != qn('w:sectPr'):
            doc._element.body.remove(child)
    doc.add_heading(config['title'], 0)
    for text in paragraphs:
        doc.add_paragraph(text)
    tables, figures = [], 0
    section = doc.sections[0]
    for item in prepared:
        if item['type'] == 'paragraph':
            doc.add_paragraph(item['text'])
            continue
        doc.add_page_break()
        doc.add_heading(item['title'], 1)
        if item['type'] == 'table':
            element = item['element']
            doc._element.body.insert_element_before(element, 'w:sectPr')
            obj = Table(element, doc)
            for tr, values in zip(element.findall(qn('w:tr')), item['rows']):
                for tc, value in zip(tr.findall(qn('w:tc')), values):
                    set_cell_text(_Cell(tc, obj), value)
            item['filled_element'] = deepcopy(element)
            tables.append(item)
        else:
            doc.add_picture(str(item['source_path']),
                            width=section.page_width - section.left_margin - section.right_margin)
            figures += 1
        for note in item.get('notes', []):
            doc.add_paragraph(note)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    verify(output, tables, source, figures)
    require(all(sha(x['source_path']) == x['sha256'] for x in prepared if 'source_path' in x),
            '来源在导出期间变化')
    require(sha(manifest) == manifest_hash and sha(reference) == reference_hash, '清单或参考Word在导出期间变化')
    receipt = {'status': 'STRUCTURE_VALUES_REFERENCE_FORMAT_PASS_VISUAL_REVIEW_PENDING',
               'tables': len(tables), 'figures': figures, 'word': str(output.resolve()),
               'word_sha256': sha(output), 'manifest_sha256': manifest_hash,
               'reference_docx': str(reference), 'reference_sha256': reference_hash,
               'xml_compatibility_fixes': compatibility_fixes,
               'regressions_rerun': False,
               'page_cm': {name: getattr(section, name).cm for name in
                           ['page_width', 'page_height', 'left_margin', 'right_margin', 'top_margin', 'bottom_margin']},
               'verification_scope': 'CSV逐格一致、页面/样式与参考稿一致、复制的表格格式完整保留、图片数；不验证估计数值、最终Stata核对或PDF视觉',
               'sources': [{'title': x['title'], 'path': str(x['source_path']), 'sha256': x['sha256'],
                            **({'reference_table': x['reference_table']} if x['type'] == 'table' else {})}
                           for x in prepared if 'source_path' in x]}
    if full_results:
        complete = copy_complete_word(output, project_root)
        receipt['complete_word_copy'] = str(complete.resolve())
        receipt['complete_word_copy_sha256'] = sha(complete)
    output.with_suffix('.核验.json').write_text(json.dumps(receipt, ensure_ascii=False, indent=2), encoding='utf-8')
    return receipt


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('manifest', type=Path)
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--full-results', action='store_true', help='全部实证结果汇编，额外存一份Word到项目根/完整实证')
    parser.add_argument('--project-root', type=Path, help='真实项目根目录；--full-results时必须提供')
    args = parser.parse_args()
    print(json.dumps(build(args.manifest, args.output, full_results=args.full_results,
                           project_root=args.project_root), ensure_ascii=False, indent=2))
